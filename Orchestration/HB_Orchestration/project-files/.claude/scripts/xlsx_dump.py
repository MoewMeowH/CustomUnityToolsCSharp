#!/usr/bin/env python3
"""/work-order 첨부 읽기 도우미 — 기획서 엑셀(xlsx/xlsm)·csv를 Claude가 읽기 좋은 Markdown으로 뽑는다.

사용:
    python .claude/scripts/xlsx_dump.py <파일> [--out review_tmp/work-order/<이름>] [--max-rows 200] [--max-cols 40]

출력 (stdout, UTF-8):
    - 시트마다 Markdown 표. 첫 열은 엑셀 행 번호, 머리글은 열 문자 (셀 위치로 인용할 수 있게).
      값이 하나도 없는 행·열은 뺀다. 병합 셀은 왼쪽 위 칸에만 값이 있고, 병합 범위는 표 아래에 적는다.
    - 셀 메모, 텍스트 상자·도형 안의 글, 이미지 목록 (저장 경로와 붙어 있는 셀).
이미지 (--out 폴더):
    - 시트에 붙은 이미지: <시트>_<셀>_<n>.<확장자>
    - 위치를 알 수 없는 이미지 (셀 안 이미지 등): 압축 안 xl/media/에서 꺼내 unplaced_<n>.<확장자>
    - emf·wmf는 그대로 저장한다 (Claude가 열 수 없을 수 있음 → 목록에 표시).
실패하면 stderr에 '[오류] …'를 쓰고 종료 코드 2. 필요한 것: Python 3.9+, openpyxl (csv만 읽을 땐 불필요).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import posixpath
import re
import sys
import warnings
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}
VIEWABLE = {"png", "jpeg", "jpg", "gif", "bmp", "webp"}


def resolve(base_dir: str, target: str) -> str:
    """관계(.rels)의 Target을 압축 안 경로로. '/xl/…'처럼 절대 경로로 적는 파일도 있다."""
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(base_dir, target))


def fail(message: str) -> None:
    print(f"[오류] {message}", file=sys.stderr)
    sys.exit(2)


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("|", "\\|").replace("\n", "<br>").strip()


def safe_name(name: str) -> str:
    return re.sub(r"[^\w\-]+", "_", name, flags=re.UNICODE).strip("_") or "sheet"


def markdown_table(rows: list[tuple[str, list[str]]], headers: list[str]) -> str:
    lines = ["| 행 | " + " | ".join(headers) + " |", "|---|" + "---|" * len(headers)]
    for label, cells in rows:
        lines.append(f"| {label} | " + " | ".join(cells) + " |")
    return "\n".join(lines)


# ── csv ──────────────────────────────────────────────────────────────

def dump_csv(path: Path, max_rows: int, max_cols: int) -> None:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        fail(f"'{path.name}': UTF-8도 CP949도 아닙니다. 인코딩을 확인해 주세요.")

    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    width = min(max((len(r) for r in rows), default=0), max_cols)

    print(f"# {path.name} (csv, {encoding}, 내용 있는 행 {len(rows)}개)\n")
    if not rows:
        print("(비어 있음)")
        return
    headers = [str(i + 1) for i in range(width)]
    body = [(str(i + 1), [cell_text(c) for c in (r + [""] * width)[:width]]) for i, r in enumerate(rows[:max_rows])]
    print(markdown_table(body, headers))
    if len(rows) > max_rows:
        print(f"\n(행 {max_rows}개까지만 표시, 나머지 {len(rows) - max_rows}개 생략 — --max-rows로 늘릴 수 있음)")


# ── xlsx 압축 안의 도형 글 · 이미지 ───────────────────────────────────────

def sheet_drawing_map(zf: zipfile.ZipFile) -> dict[str, str]:
    """시트 이름 → drawing XML 경로."""
    names = set(zf.namelist())
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    wb_rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {r.get("Id"): r.get("Target") for r in wb_rels.findall("rel:Relationship", NS)}
    result = {}
    for sheet in wb.findall("m:sheets/m:sheet", NS):
        target = rid_to_target.get(sheet.get(f"{{{NS['r']}}}id"), "")
        sheet_path = resolve("xl", target)
        rels_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels")
        if rels_path not in names:
            continue
        for rel in ET.fromstring(zf.read(rels_path)).findall("rel:Relationship", NS):
            if rel.get("Type", "").endswith("/drawing"):
                result[sheet.get("name")] = resolve(posixpath.dirname(sheet_path), rel.get("Target"))
    return result


def anchor_cell(anchor: ET.Element) -> str:
    frm = anchor.find("xdr:from", NS)
    if frm is None:
        return "위치 모름"
    from openpyxl.utils import get_column_letter
    col = int(frm.findtext("xdr:col", "0", NS)) + 1
    row = int(frm.findtext("xdr:row", "0", NS)) + 1
    return f"{get_column_letter(col)}{row}"


def drawing_texts(zf: zipfile.ZipFile, drawing_path: str) -> list[tuple[str, str]]:
    """텍스트 상자·도형 안의 글 (붙어 있는 셀, 글)."""
    texts = []
    root = ET.fromstring(zf.read(drawing_path))
    for anchor in list(root):
        for body in anchor.iter(f"{{{NS['xdr']}}}txBody"):
            paragraphs = ["".join(t.text or "" for t in p.iter(f"{{{NS['a']}}}t")) for p in body.findall("a:p", NS)]
            text = "\n".join(p for p in paragraphs if p.strip())
            if text.strip():
                texts.append((anchor_cell(anchor), text))
    return texts


def drawing_images(zf: zipfile.ZipFile, drawing_path: str) -> list[tuple[str, str]]:
    """시트에 붙은 이미지 (붙어 있는 셀, 압축 안 경로)."""
    rels_path = posixpath.join(posixpath.dirname(drawing_path), "_rels", posixpath.basename(drawing_path) + ".rels")
    if rels_path not in zf.namelist():
        return []
    rid_to_media = {
        r.get("Id"): resolve(posixpath.dirname(drawing_path), r.get("Target"))
        for r in ET.fromstring(zf.read(rels_path)).findall("rel:Relationship", NS)
        if r.get("Type", "").endswith("/image")
    }
    images = []
    for anchor in list(ET.fromstring(zf.read(drawing_path))):
        for blip in anchor.iter(f"{{{NS['a']}}}blip"):
            media = rid_to_media.get(blip.get(f"{{{NS['r']}}}embed"))
            if media:
                images.append((anchor_cell(anchor), media))
    return images


def save_media(zf: zipfile.ZipFile, member: str, out_dir: Path, stem: str, saved: dict[str, Path]) -> tuple[Path, bool]:
    data = zf.read(member)
    digest = hashlib.sha256(data).hexdigest()
    ext = posixpath.splitext(member)[1].lower().lstrip(".") or "bin"
    target = out_dir / f"{stem}.{ext}"
    target.write_bytes(data)
    saved.setdefault(digest, target)
    return target, ext in VIEWABLE


# ── xlsx ─────────────────────────────────────────────────────────────

def dump_xlsx(path: Path, out_dir: Path, max_rows: int, max_cols: int) -> None:
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        fail("openpyxl이 없습니다. 'python -m pip install openpyxl'로 설치하거나 csv로 저장해 주세요.")

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")   # 도형 미지원 경고 — 도형 글은 아래에서 압축을 직접 읽는다
            wb = openpyxl.load_workbook(path, data_only=True)        # 수식은 마지막으로 계산된 값
            formulas = openpyxl.load_workbook(path, data_only=False)  # 계산 값이 저장되지 않은 수식 표시용
        zf = zipfile.ZipFile(path)
    except Exception as ex:  # 손상·암호·형식 오류
        fail(f"'{path.name}'을(를) 열 수 없습니다: {type(ex).__name__}: {ex}")

    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        drawings = sheet_drawing_map(zf)
    except (KeyError, ET.ParseError):
        drawings = {}
    saved: dict[str, Path] = {}   # sha256 → 저장 경로 (위치 모름 이미지와 중복 제거용)

    print(f"# {path.name} (시트 {len(wb.sheetnames)}개: {', '.join(wb.sheetnames)})")
    print(f"이미지 폴더: {out_dir.as_posix()}\n")

    for ws in wb.worksheets:
        fs = formulas[ws.title]

        def value_at(r: int, c: int) -> str:
            value = ws.cell(r, c).value
            if value is None and isinstance(fs.cell(r, c).value, str) and fs.cell(r, c).value.startswith("="):
                return cell_text(fs.cell(r, c).value) + " (계산 값 없음)"
            return cell_text(value)

        hidden = "" if ws.sheet_state == "visible" else f" ({ws.sheet_state})"
        print(f"## 시트: {ws.title}{hidden}\n")

        filled = [(c.row, c.column) for row in fs.iter_rows() for c in row if c.value is not None and str(c.value).strip()]
        if filled:
            rows = sorted({r for r, _ in filled})
            cols = sorted({c for _, c in filled})
            shown_cols = cols[:max_cols]
            headers = [get_column_letter(c) for c in shown_cols]
            body = [(str(r), [value_at(r, c) for c in shown_cols]) for r in rows[:max_rows]]
            print(markdown_table(body, headers))
            if len(rows) > max_rows:
                print(f"\n(내용 있는 행 {len(rows)}개 중 {max_rows}개만 표시 — --max-rows로 늘릴 수 있음)")
            if len(cols) > max_cols:
                print(f"\n(내용 있는 열 {len(cols)}개 중 {max_cols}개만 표시 — --max-cols로 늘릴 수 있음)")
        else:
            print("(셀 값 없음)")

        merged = [str(m) for m in ws.merged_cells.ranges]
        if merged:
            print(f"\n병합 셀: {', '.join(merged[:50])}" + (" …" if len(merged) > 50 else ""))

        comments = [(c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment]
        if comments:
            print("\n메모:")
            for coord, text in comments:
                print(f"- {coord}: {cell_text(text)}")

        drawing = drawings.get(ws.title)
        if drawing:
            try:
                texts = drawing_texts(zf, drawing)
                images = drawing_images(zf, drawing)
            except (KeyError, ET.ParseError) as ex:
                texts, images = [], []
                print(f"\n[주의] 도형·이미지를 읽지 못했습니다: {type(ex).__name__}")
            if texts:
                print("\n텍스트 상자·도형:")
                for cell, text in texts:
                    print(f"- {cell}: {cell_text(text)}")
            if images:
                print("\n이미지:")
                for n, (cell, member) in enumerate(images, 1):
                    target, viewable = save_media(zf, member, out_dir, f"{safe_name(ws.title)}_{cell}_{n}", saved)
                    note = "" if viewable else " (열 수 없는 형식일 수 있음)"
                    print(f"- {cell}: {target.as_posix()}{note}")
        print()

    # 시트에 붙지 않은 이미지 (셀 안 이미지, 머리글 이미지 등)
    leftovers = []
    for member in sorted(n for n in zf.namelist() if n.startswith("xl/media/")):
        if hashlib.sha256(zf.read(member)).hexdigest() in saved:
            continue
        target, viewable = save_media(zf, member, out_dir, f"unplaced_{len(leftovers) + 1}", saved)
        leftovers.append((target, viewable))
    if leftovers:
        print("## 위치를 알 수 없는 이미지 (셀 안 이미지 등)\n")
        for target, viewable in leftovers:
            print(f"- {target.as_posix()}" + ("" if viewable else " (열 수 없는 형식일 수 있음)"))


def main() -> None:
    parser = argparse.ArgumentParser(description="기획서 xlsx/csv를 Markdown으로 뽑는다 (/work-order 첨부 읽기).")
    parser.add_argument("file")
    parser.add_argument("--out", help="이미지 저장 폴더 (기본: review_tmp/work-order/<파일 이름>)")
    parser.add_argument("--max-rows", type=int, default=200)
    parser.add_argument("--max-cols", type=int, default=40)
    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")   # Windows 콘솔 기본 인코딩(cp949) 회피

    path = Path(args.file)
    if not path.is_file():
        fail(f"'{args.file}' 파일이 없습니다.")
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        dump_csv(path, args.max_rows, args.max_cols)
    elif suffix in (".xlsx", ".xlsm"):
        out_dir = Path(args.out) if args.out else Path("review_tmp/work-order") / safe_name(path.stem)
        dump_xlsx(path, out_dir, args.max_rows, args.max_cols)
    elif suffix == ".xls":
        fail("예전 형식(.xls)은 읽지 못합니다. 엑셀에서 .xlsx로 다시 저장해 주세요.")
    else:
        fail(f"지원하지 않는 형식 '{suffix}' (xlsx, xlsm, csv만)")


if __name__ == "__main__":
    main()
